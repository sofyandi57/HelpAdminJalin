"""
OCR PDF PO -> Excel/CSV, versi ringan (Tkinter, tanpa Streamlit/pandas).
Dibuat khusus supaya bisa dibungkus jadi .exe berukuran kecil (target <50MB).
Tesseract OCR TIDAK dibundel di sini -- harus sudah terpasang terpisah di sistem.
"""

import csv
import glob
import os
import queue
import sys
import threading
import tkinter as tk
from tkinter import filedialog, messagebox, ttk

from openpyxl import Workbook

from ocr_utils import configure_tesseract
from po_extract import COLUMNS, extract_po

if getattr(sys, "frozen", False):
    BASE_DIR = os.path.dirname(sys.executable)
else:
    BASE_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_DIR = os.path.join(BASE_DIR, "output")


class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("OCR PDF PO -> Excel/CSV")
        self.geometry("1000x650")

        self.pdf_paths = []  # list of (nama_tampilan, path_asli)
        self.result_rows = []
        self.worker_queue = queue.Queue()

        self._build_widgets()
        self._check_tesseract()

    # ------------------------------------------------------------------ UI
    def _build_widgets(self):
        top = ttk.Frame(self, padding=10)
        top.pack(fill="x")

        self.tesseract_label = ttk.Label(top, text="Mengecek Tesseract OCR...", foreground="gray")
        self.tesseract_label.pack(anchor="w")

        btn_row = ttk.Frame(self, padding=(10, 0))
        btn_row.pack(fill="x")
        ttk.Button(btn_row, text="Tambah File PDF...", command=self.add_files).pack(side="left", padx=(0, 6))
        ttk.Button(btn_row, text="Tambah dari Folder...", command=self.add_folder).pack(side="left", padx=(0, 6))
        ttk.Button(btn_row, text="Bersihkan Daftar", command=self.clear_files).pack(side="left", padx=(0, 6))

        list_frame = ttk.LabelFrame(self, text="File PDF yang akan diproses", padding=5)
        list_frame.pack(fill="x", padx=10, pady=8)
        self.file_listbox = tk.Listbox(list_frame, height=5)
        self.file_listbox.pack(fill="x")

        process_row = ttk.Frame(self, padding=(10, 0))
        process_row.pack(fill="x")
        self.process_btn = ttk.Button(process_row, text="Proses OCR", command=self.start_processing)
        self.process_btn.pack(side="left")
        self.progress = ttk.Progressbar(process_row, mode="determinate")
        self.progress.pack(side="left", fill="x", expand=True, padx=10)
        self.status_label = ttk.Label(process_row, text="")
        self.status_label.pack(side="left")

        result_frame = ttk.LabelFrame(self, text="Hasil", padding=5)
        result_frame.pack(fill="both", expand=True, padx=10, pady=8)

        tree_scroll_y = ttk.Scrollbar(result_frame, orient="vertical")
        tree_scroll_x = ttk.Scrollbar(result_frame, orient="horizontal")
        self.tree = ttk.Treeview(
            result_frame,
            columns=COLUMNS,
            show="headings",
            yscrollcommand=tree_scroll_y.set,
            xscrollcommand=tree_scroll_x.set,
        )
        for col in COLUMNS:
            self.tree.heading(col, text=col)
            self.tree.column(col, width=110, stretch=False)
        tree_scroll_y.config(command=self.tree.yview)
        tree_scroll_x.config(command=self.tree.xview)
        self.tree.grid(row=0, column=0, sticky="nsew")
        tree_scroll_y.grid(row=0, column=1, sticky="ns")
        tree_scroll_x.grid(row=1, column=0, sticky="ew")
        result_frame.rowconfigure(0, weight=1)
        result_frame.columnconfigure(0, weight=1)

        export_row = ttk.Frame(self, padding=10)
        export_row.pack(fill="x")
        self.save_csv_btn = ttk.Button(export_row, text="Simpan sebagai CSV", command=self.save_csv, state="disabled")
        self.save_csv_btn.pack(side="left", padx=(0, 6))
        self.save_xlsx_btn = ttk.Button(export_row, text="Simpan sebagai Excel", command=self.save_xlsx, state="disabled")
        self.save_xlsx_btn.pack(side="left", padx=(0, 6))
        self.cumulative_btn = ttk.Button(
            export_row, text="Tambahkan ke dataset kumulatif", command=self.save_cumulative, state="disabled"
        )
        self.cumulative_btn.pack(side="left")

    def _check_tesseract(self):
        cmd = configure_tesseract()
        if cmd:
            self.tesseract_label.config(text=f"Tesseract OCR terdeteksi: {cmd}", foreground="green")
        else:
            self.tesseract_label.config(
                text="Tesseract OCR TIDAK ditemukan di komputer ini. Install dulu (lihat README.md).",
                foreground="red",
            )

    # ------------------------------------------------------------- actions
    def add_files(self):
        paths = filedialog.askopenfilenames(title="Pilih file PDF", filetypes=[("PDF", "*.pdf")])
        for p in paths:
            self.pdf_paths.append((os.path.basename(p), p))
            self.file_listbox.insert("end", os.path.basename(p))

    def add_folder(self):
        folder = filedialog.askdirectory(title="Pilih folder berisi PDF")
        if not folder:
            return
        found = sorted(glob.glob(os.path.join(folder, "**", "*.pdf"), recursive=True))
        for p in found:
            name = os.path.relpath(p, folder)
            self.pdf_paths.append((name, p))
            self.file_listbox.insert("end", name)
        if not found:
            messagebox.showinfo("Info", "Tidak ada file PDF ditemukan di folder tersebut.")

    def clear_files(self):
        self.pdf_paths.clear()
        self.file_listbox.delete(0, "end")

    def start_processing(self):
        if not self.pdf_paths:
            messagebox.showwarning("Peringatan", "Belum ada file PDF yang ditambahkan.")
            return
        self.process_btn.config(state="disabled")
        self.save_csv_btn.config(state="disabled")
        self.save_xlsx_btn.config(state="disabled")
        self.cumulative_btn.config(state="disabled")
        for row in self.tree.get_children():
            self.tree.delete(row)
        self.progress.config(maximum=len(self.pdf_paths), value=0)

        thread = threading.Thread(target=self._process_worker, daemon=True)
        thread.start()
        self.after(100, self._poll_queue)

    def _process_worker(self):
        rows = []
        for name, path in self.pdf_paths:
            self.worker_queue.put(("status", f"Memproses: {name}"))
            try:
                rows.extend(extract_po(path, name))
            except Exception as exc:  # noqa: BLE001
                rows.append({**{c: "" for c in COLUMNS}, "Nama File": name, "Error": str(exc)})
            self.worker_queue.put(("progress", 1))
        self.worker_queue.put(("done", rows))

    def _poll_queue(self):
        try:
            while True:
                kind, payload = self.worker_queue.get_nowait()
                if kind == "status":
                    self.status_label.config(text=payload)
                elif kind == "progress":
                    self.progress.step(payload)
                elif kind == "done":
                    self._on_processing_done(payload)
                    return
        except queue.Empty:
            pass
        self.after(100, self._poll_queue)

    def _on_processing_done(self, rows):
        self.result_rows = rows
        for row in rows:
            self.tree.insert("", "end", values=[row.get(c, "") for c in COLUMNS])
        self.status_label.config(text=f"Selesai. {len(rows)} baris.")
        self.process_btn.config(state="normal")
        has_rows = bool(rows)
        self.save_csv_btn.config(state="normal" if has_rows else "disabled")
        self.save_xlsx_btn.config(state="normal" if has_rows else "disabled")
        self.cumulative_btn.config(state="normal" if has_rows else "disabled")
        n_error = sum(1 for r in rows if r.get("Error"))
        if n_error:
            messagebox.showwarning("Selesai dengan peringatan", f"{n_error} file gagal diproses sepenuhnya — lihat kolom Error.")

    # ------------------------------------------------------------- export
    def save_csv(self):
        path = filedialog.asksaveasfilename(defaultextension=".csv", filetypes=[("CSV", "*.csv")], initialfile="hasil_ocr.csv")
        if not path:
            return
        write_csv(path, self.result_rows)
        messagebox.showinfo("Tersimpan", f"Hasil disimpan ke:\n{path}")

    def save_xlsx(self):
        path = filedialog.asksaveasfilename(defaultextension=".xlsx", filetypes=[("Excel", "*.xlsx")], initialfile="hasil_ocr.xlsx")
        if not path:
            return
        write_xlsx(path, self.result_rows)
        messagebox.showinfo("Tersimpan", f"Hasil disimpan ke:\n{path}")

    def save_cumulative(self):
        os.makedirs(OUTPUT_DIR, exist_ok=True)
        xlsx_path = os.path.join(OUTPUT_DIR, "hasil_ocr.xlsx")
        csv_path = os.path.join(OUTPUT_DIR, "hasil_ocr.csv")
        existing = read_xlsx(xlsx_path) if os.path.isfile(xlsx_path) else []
        combined = existing + self.result_rows
        write_xlsx(xlsx_path, combined)
        write_csv(csv_path, combined)
        messagebox.showinfo("Tersimpan", f"Total baris di dataset kumulatif: {len(combined)}\n({xlsx_path})")


def write_csv(path, rows):
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=COLUMNS)
        writer.writeheader()
        for row in rows:
            writer.writerow({c: row.get(c, "") for c in COLUMNS})


def write_xlsx(path, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "Hasil OCR"
    ws.append(COLUMNS)
    for row in rows:
        ws.append([row.get(c, "") for c in COLUMNS])
    wb.save(path)


def read_xlsx(path):
    from openpyxl import load_workbook

    wb = load_workbook(path)
    ws = wb.active
    header = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    rows = []
    for excel_row in ws.iter_rows(min_row=2, values_only=True):
        rows.append({header[i]: (excel_row[i] if excel_row[i] is not None else "") for i in range(len(header))})
    return rows


if __name__ == "__main__":
    App().mainloop()
